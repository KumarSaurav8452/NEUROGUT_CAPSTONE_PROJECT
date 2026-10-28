"""FR-3: map each run to group in {Healthy, SCZ, Exclude} using config/label_maps/<StudyID>.yaml.

A label map is an ordered list of rules; the first rule that matches a run wins.

  - field / match / group [/ reason]     regex (case-insensitive) on a raw-metadata column
  - supplement: {...}                    join a paper's supplementary table on a key column and
                                         take the group (plus age/sex/BMI/published read count) from it;
                                         optional `filter` (keep rows whose column equals a value) and
                                         `key_transform` (regex + format + offset, when the paper renames
                                         samples, e.g. Xu: CNGB wmb-63 -> paper SCZ01)
  default: {group: Exclude, reason: ...} for runs no rule matched
"""

from __future__ import annotations

import re
from pathlib import Path

import openpyxl
import pandas as pd
import yaml

from scz_data.settings import Settings, get_logger

GROUPS = {"Healthy", "SCZ", "Exclude"}
LABEL_COLUMNS = ["group", "label_rule", "exclusion_reason", "age", "sex", "bmi", "diagnosis_detail",
                 "published_reads"]


def _norm(header) -> str:
    return " ".join(str(header).split()) if header is not None else ""


def load_supplement(spec: dict, root: Path) -> pd.DataFrame:
    """Read the rows of one table from a supplementary xlsx. Header names are whitespace-normalised."""
    ws = openpyxl.load_workbook(root / spec["file"], data_only=True)
    ws = ws[spec["sheet"]] if spec.get("sheet") else ws.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    header_i = spec["header_row"] - 1
    last = spec.get("last_row") or len(rows)
    header = [_norm(h) for h in rows[header_i]]
    # Some supplements repeat a column name (e.g. Zhu Supp Data 1 has 'BGI ID' twice); keep the first.
    keep = [i for i, h in enumerate(header) if h and header.index(h) == i]
    data = [[r[i] for i in keep] for r in rows[header_i + 1 : last] if r[keep[0]] not in (None, "")]
    return pd.DataFrame(data, columns=[header[i] for i in keep])


def _supplement_labels(spec: dict, root: Path) -> pd.DataFrame:
    sup = load_supplement(spec, root)
    for col, value in (spec.get("filter") or {}).items():
        sup = sup[sup[_norm(col)].astype(str).str.strip() == str(value)]
    key = _norm(spec["key"])
    if sup[key].duplicated().any():
        raise ValueError(f"{spec['file']}: duplicate keys in {key!r}")
    out = pd.DataFrame({"_key": sup[key].astype(str).str.strip()})
    raw_group = sup[_norm(spec["group_column"])].astype(str).str.strip()
    out["group"] = raw_group.map(spec["groups"])
    out["_raw_group"] = raw_group
    for target, col in spec.get("columns", {}).items():
        out[target] = sup[_norm(col)]
    if "bmi" in out:
        out["bmi"] = pd.to_numeric(out["bmi"], errors="coerce")  # '#VALUE!' / 'NA' -> NaN
    for num in ("age", "published_reads"):
        if num in out:
            out[num] = pd.to_numeric(out[num], errors="coerce")
    if "sex" in out and spec.get("sex_codes"):
        codes = {str(k): v for k, v in spec["sex_codes"].items()}
        out["sex"] = out["sex"].map(lambda v: codes.get(str(v).strip(), pd.NA) if pd.notna(v) else pd.NA)
    return out.set_index("_key")


def transform_keys(keys: pd.Series, spec: dict | None) -> pd.Series:
    """`key_transform: {regex, format, offset}`: e.g. regex '^wmb-([0-9]+)$', format 'SCZ{:02d}', offset -62."""
    if not spec:
        return keys
    pattern = re.compile(spec["regex"])

    def one(v: str):
        m = pattern.match(v)
        return spec["format"].format(int(m.group(1)) + int(spec.get("offset", 0))) if m else v

    return keys.map(one)


def apply_label_map(df: pd.DataFrame, label_map: dict, root: Path) -> tuple[pd.DataFrame, list[str]]:
    """Return df with LABEL_COLUMNS added, plus a list of warnings."""
    df = df.copy()
    for col in LABEL_COLUMNS:
        df[col] = pd.NA
    warnings = []
    unassigned = pd.Series(True, index=df.index)

    for i, rule in enumerate(label_map["rules"], start=1):
        if "supplement" in rule:
            spec = rule["supplement"]
            sup = _supplement_labels(spec, root)
            keys = transform_keys(df[spec["join_on"]].astype(str).str.strip(), spec.get("key_transform"))
            hit = unassigned & keys.isin(sup.index)
            matched = sup.loc[keys[hit]]
            bad = matched["group"].isna()
            if bad.any():
                warnings.append(f"rule {i}: unmapped supplement groups {sorted(matched.loc[bad, '_raw_group'].unique())}")
            for col in ("group", "age", "sex", "bmi", "diagnosis_detail", "published_reads"):
                if col in matched:
                    df.loc[hit, col] = matched[col].values
            df.loc[hit, "label_rule"] = f"{i}:supplement:{Path(spec['file']).name}"
            orphans = set(sup.index) - set(keys)
            if orphans:
                warnings.append(f"rule {i}: {len(orphans)} supplement rows have no run in metadata, e.g. {sorted(orphans)[:3]}")
        else:
            pattern = re.compile(rule["match"], re.IGNORECASE)
            values = df[rule["field"]].fillna("").astype(str)
            hit = unassigned & values.map(lambda v: bool(pattern.search(v)))
            df.loc[hit, "group"] = rule["group"]
            df.loc[hit, "label_rule"] = f"{i}:{rule['field']}~/{rule['match']}/"
            if rule["group"] == "Exclude":
                df.loc[hit, "exclusion_reason"] = rule.get("reason", f"matched {rule['match']}")
        unassigned &= df["group"].isna()

    default = label_map.get("default", {"group": "Exclude", "reason": "no label rule matched"})
    df.loc[unassigned, "group"] = default["group"]
    df.loc[unassigned, "label_rule"] = "default"
    if default["group"] == "Exclude":
        df.loc[unassigned, "exclusion_reason"] = default.get("reason", "no label rule matched")
    if unassigned.any():
        warnings.append(f"{int(unassigned.sum())} runs fell through to default ({default['group']})")

    unknown = set(df["group"].dropna()) - GROUPS
    if unknown:
        raise ValueError(f"label map produced unknown groups {unknown}")
    return df, warnings


def sample_counts(df: pd.DataFrame) -> dict[str, int]:
    """Unique samples per group (runs of one sample are merged later, in FR-4)."""
    per_sample = df.groupby("sample_accession")["group"].agg(lambda g: g.iloc[0] if g.nunique() == 1 else "CONFLICT")
    return per_sample.value_counts().to_dict()


def run(settings: Settings, studies: list[dict]) -> dict[str, str]:
    from scz_data.validate import compare_counts

    out_dir = settings.path("labelled")
    out_dir.mkdir(parents=True, exist_ok=True)
    status, summary = {}, []
    tol = settings.cfg["filters"]["count_tolerance"]
    for study in studies:
        sid = study["study_id"]
        log = get_logger(settings, "label", sid)
        raw = settings.path("raw_metadata") / f"{sid}.tsv"
        map_path = settings.path("label_maps") / f"{sid}.yaml"
        if not raw.exists():
            log.warning(f"SKIPPED: no raw metadata ({raw.name})")
            status[sid] = "SKIPPED: no raw metadata"
            continue
        try:
            with open(map_path, encoding="utf-8") as fh:
                label_map = yaml.safe_load(fh)
            df = pd.read_csv(raw, sep="\t", dtype=str, keep_default_na=False, na_values=[""])
            labelled, warnings = apply_label_map(df, label_map, settings.root)
            for w in warnings:
                log.warning(w)
            counts = sample_counts(labelled)
            if counts.get("CONFLICT"):
                raise RuntimeError(f"{counts['CONFLICT']} samples have runs with conflicting labels")
            out = out_dir / f"{sid}.tsv"
            labelled.to_csv(out, sep="\t", index=False)
            exp = study["expected"]
            for group, key in (("Healthy", "healthy_raw"), ("SCZ", "scz_raw")):
                got = counts.get(group, 0)
                summary.append({"study_id": sid, "group": group, "expected_raw": exp[key], "labelled": got,
                                "status": compare_counts(exp[key], got, tol)})
            summary.append({"study_id": sid, "group": "Exclude", "expected_raw": pd.NA,
                            "labelled": counts.get("Exclude", 0), "status": ""})
            log.info(f"wrote {out.relative_to(settings.root)}: " + ", ".join(f"{k}={v}" for k, v in sorted(counts.items())))
            status[sid] = "OK"
        except Exception as e:
            log.exception(f"FAILED: {e}")
            status[sid] = f"FAILED: {e}"
    if summary:
        res = settings.path("results")
        res.mkdir(parents=True, exist_ok=True)
        table = pd.DataFrame(summary)
        table.to_csv(res / "label_counts_vs_table_s1.tsv", sep="\t", index=False)
        print("\nLabelled samples vs Table S1 raw counts (unique samples, before FR-4 filtering):")
        print(table.to_string(index=False))
    return status
